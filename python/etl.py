import pandas as pd
from datetime import date, timedelta, datetime
import os
import csv
import psycopg2
from dotenv import load_dotenv
import paramiko
import json
from openpyxl import load_workbook

import warnings
warnings.simplefilter('ignore', UserWarning)

def sftp_move_files(rasp_ip, rasp_ssh_port, usernm_sftp, rasp_pwd, win_path_plan_gas, rasp_home_path):

    try:
        transport = paramiko.Transport((rasp_ip, int(rasp_ssh_port)))
        transport.connect(username=usernm_sftp, password=rasp_pwd)
        
        sftp = paramiko.SFTPClient.from_transport(transport)
        
        # Upload folder            
        for root, dirs, files in os.walk(win_path_plan_gas):
            remote_path = os.path.join(rasp_home_path, os.path.relpath(root, win_path_plan_gas)).replace('\\', '/')
            try:
                sftp.mkdir(remote_path)
                print(f'Directory created: {remote_path}')
            except IOError:
                # Directory already exists
                pass
            
            for file in files:
                local_file = os.path.join(root, file)
                remote_file = os.path.join(remote_path, file).replace('\\', '/')
                sftp.put(local_file, remote_file)
                # print(f'File uploaded: {local_file} to {remote_file}')
        
        print('SFTP ran successfully')
        # Close connection
        sftp.close()
        transport.close()

    except Exception as e:
        print(f'An error occurred: {str(e)}')
    except paramiko.AuthenticationException:
        print('Authentication failed, please verify your credentials.')
    except paramiko.SSHException as ssh_exception:
        print(f'Unable to establish SSH connection: {ssh_exception}')
    except Exception as e:
        print(f'Operation error: {e}')

####

def valor_empty(path_base, dataframe):
    try:
        if dataframe['valor'].isnull().any():

            dataframe.to_csv(path_base + '/log_file.csv', index = False, sep = ';')

            raise ValueError(f"The column 'valor' contains empty values.")
    except KeyError:
        raise KeyError(f"The column 'valor' does not exist in the DataFrame.")
    except Exception as e:
        raise e

def fato_2023(path_base, folder_tabelas):

    df_fato_plano_2023 = pd.read_excel(path_base + folder_tabelas + 'fato_plano_gastos_2023.xlsx')

    df_processed_2023 = df_fato_plano_2023.dropna(axis = 'index', subset = ['data'], how = 'all', inplace = False).fillna(value = {'quantidade': 0, 'preco': 0, 'valor': 0})

    df_processed_2023 = df_processed_2023.reset_index()
    df_processed_2023 = df_processed_2023.rename(columns = {'index': 'id_fato_plano_gastos'})
    df_processed_2023['id_fato_plano_gastos'] = df_processed_2023.index + 1

    df_processed_2023['quantidade'] = df_processed_2023['quantidade'].str.replace(',', '.')
    df_processed_2023['preco'] = df_processed_2023['preco'].str.replace(',', '.')
    df_processed_2023['valor'] = df_processed_2023['valor'].str.replace(',', '.')

    df_processed_2023.astype({'id_fato_plano_gastos': 'int32', 'quantidade': 'float'})
    df_processed_2023['data'] = pd.to_datetime(df_processed_2023['data'], format = '%Y-%m-%d')

    valor_empty(path_base, df_processed_2023)

    df_processed_2023.to_csv(path_base + folder_tabelas + 'processed/fato_plan_gas_proc_2023.csv', index = False, sep = ';')

    return df_processed_2023

def fato_2024(path_base, folder_tabelas):

    df_fato_plano_2024 = pd.read_excel(path_base + folder_tabelas + 'fato_plano_gastos_2024.xlsx')

    df_processed_2024 = df_fato_plano_2024.dropna(axis = 'index', subset = ['data'], how = 'all', inplace = False).fillna(value = {'quantidade': 0, 'preco': 0, 'valor': 0})

    df_processed_2024 = df_processed_2024.reset_index()
    df_processed_2024 = df_processed_2024.rename(columns = {'index': 'id_fato_plano_gastos'})
    df_processed_2024['id_fato_plano_gastos'] = df_processed_2024.index + 1

    df_processed_2024['quantidade'] = df_processed_2024['quantidade'].str.replace(',', '.')
    df_processed_2024['preco'] = df_processed_2024['preco'].str.replace(',', '.')
    df_processed_2024['valor'] = df_processed_2024['valor'].str.replace(',', '.')

    df_processed_2024.astype({'id_fato_plano_gastos': 'int32', 'quantidade': 'float'})
    df_processed_2024['data'] = pd.to_datetime(df_processed_2024['data'], format = '%Y-%m-%d')

    valor_empty(path_base, df_processed_2024)

    df_processed_2024.to_csv(path_base + folder_tabelas + 'processed/fato_plan_gas_proc_2024.csv', index = False, sep = ';')

    return df_processed_2024

def join_anos(path_base, folder_tabelas, df_processed_2023, df_processed_2024):

    df_join_anos = pd.concat([df_processed_2023, df_processed_2024], axis = 0)

    # print(df_join_anos)
    valor_empty(path_base, df_join_anos)

    df_join_anos.to_csv(path_base + folder_tabelas + 'processed/fato_plan_gas_concat_anos.csv', index = False, sep = ';')

def notas_fiscais_2024(path_base, folder_tabelas):

    FOLDER_NOTAS_FISCAIS = 'notas_fiscais/'
    df_notas_fiscais_2024 = pd.read_excel(path_base + 'notas_fiscais_2024.xlsx')

    df_notas_fiscais_2024 = df_notas_fiscais_2024.dropna(axis = 'index', subset = ['cod_item'], how = 'all', inplace = False)

    df_notas_fiscais_2024 = df_notas_fiscais_2024.dropna(axis = 'index', subset = ['data'], how = 'all', inplace = False)

    df_notas_fiscais_2024 = df_notas_fiscais_2024.reset_index()
    df_notas_fiscais_2024 = df_notas_fiscais_2024.rename(columns = {'index': 'id_notas_fiscais_gastos'})

    df_notas_fiscais_2024['id_notas_fiscais_gastos'] = df_notas_fiscais_2024.index + 1
    # print(df_notas_fiscais_2024.dtypes)

    df_notas_fiscais_2024['cod_item'] = df_notas_fiscais_2024['cod_item'].apply(int)
    df_notas_fiscais_2024['contador_itens'] = df_notas_fiscais_2024['contador_itens'].apply(int)
    
    df_notas_fiscais_2024.astype(
        {
            'id_notas_fiscais_gastos': 'int32',
            'qtd_itens': 'float',
            'valor_tributo' : 'float', 
            'valor_item': 'float',
            'categoria': 'object'
        }
    )
    
    # print(df_notas_fiscais_2024.dtypes)

    df_notas_fiscais_2024['data'] = pd.to_datetime(df_notas_fiscais_2024['data'], format = '%Y-%m-%d')
    df_notas_fiscais_2024['categoria'] = df_notas_fiscais_2024['categoria'].str.strip()

    # print(df_notas_fiscais_2024)

    df_notas_fiscais_2024.to_csv(path_base + folder_tabelas + 'processed/fato_notas_2024.csv', index = False, sep = ';')

def atv_pro(path_base, folder_tabelas, product_to_cat, product_to_sector, atv_pro_load_years):

    df_join_atv_pro = pd.DataFrame()
    
    cat = json.loads(product_to_cat)
    sec = json.loads(product_to_sector)
    # print(cat)
    # print(sec)

    for i in atv_pro_load_years: 

        folder_atv_pro = 'bov\\' + i +  '\\' + 'atv_pro\\'

        df_atv = pd.read_excel(path_base + folder_atv_pro + 'atv_' + i + '.xlsx')
        df_pro = pd.read_excel(path_base + folder_atv_pro + 'pro_' + i + '.xlsx')

        dataset_cols = {
            'Data': 'data',
            'Movimentação': 'movimentacao',
            'Produto': 'produto',
            'Quantidade': 'qtd',
            'Preço unitário': 'preco_unit',
            'Valor da Operação': 'valor_operacao',
        }

        df_atv = df_atv.rename(columns = dataset_cols)
        df_pro = df_pro.rename(columns = dataset_cols)

        df_atv_processed = pd.DataFrame()
        df_pro_processed = pd.DataFrame()

        def find_key(value, mapping):
            
            for key, values in mapping.items():
                if value in values:
                    return key
            return None
        
        if i != atv_pro_load_years[1]: # this year doesn't have atv
            df_atv_processed['id'] = df_atv.index + 1
            df_atv_processed['data'] = pd.to_datetime(df_atv['data'], format = '%d/%m/%Y')
            df_atv_processed['movimentacao'] = df_atv['movimentacao']
            df_atv_processed['produto_nm'] = df_atv['produto']
            df_atv_processed[['produto', 'produto_nm']] = df_atv['produto'].str.split(n = 1, pat = ' - ', expand = True)
            df_atv_processed['sector'] = df_atv_processed['produto'].apply(lambda x: find_key(x, cat))
            df_atv_processed['category'] = df_atv_processed['produto'].apply(lambda x: find_key(x, sec))
            df_atv_processed['qtd'] = df_atv['qtd'].apply(int)
            df_atv_processed['preco_unit'] = df_atv['preco_unit']
            df_atv_processed['valor_operacao'] = df_atv['valor_operacao']
            # print(df_atv_processed.head())

        df_pro_processed['id'] = df_pro.index + 1
        df_pro_processed['data'] = pd.to_datetime(df_pro['data'], format = '%d/%m/%Y')#'%Y-%m-%d
        df_pro_processed['movimentacao'] = df_pro['movimentacao']
        df_pro_processed['produto_nm'] = df_pro['produto']
        df_pro_processed[['produto', 'produto_nm']] = df_pro['produto'].str.split(n = 1, pat = ' - ', expand = True)     
        df_pro_processed['sector'] = df_pro_processed['produto'].apply(lambda x: find_key(x, cat))
        df_pro_processed['category'] = df_pro_processed['produto'].apply(lambda x: find_key(x, sec))
        df_pro_processed['qtd'] = df_pro['qtd'].apply(int)
        df_pro_processed['preco_unit'] = df_pro['preco_unit'].round(3)
        df_pro_processed['valor_operacao'] = df_pro['valor_operacao']
        # print(df_pro_processed.head())
        
        df_join_atv_pro = pd.concat([df_join_atv_pro, df_atv_processed, df_pro_processed], axis = 0)

    # print(df_join_atv_pro)
    df_join_atv_pro.to_csv(path_base + folder_tabelas + 'processed/fato_atv_pro_concat.csv', index = False, sep = ';')

def nu_pf(path_base, folder_tabelas):
        
    folder_tabelas = 'nubank_pf\\2024\\'

    df_nu_pf = pd.DataFrame(columns = ['date', 'category', 'title', 'amount'])
    df_proc_nu_pf = pd.DataFrame(columns = ['data',	'in_out', 'categoria', 'investimento', 'quantidade', 'preco', 'valor', 'status', 'descricao'])
    
    current_month = datetime.now().month
    # print(f'The current month is: {current_month}')

    for i in range(6, current_month, 1):
        try:
            df_fatura_mes_i = pd.read_csv(path_base + folder_tabelas + 'nubank-2024-' + str(i).zfill(2) + '.csv')

            df_nu_pf = pd.concat([df_nu_pf, df_fatura_mes_i], ignore_index = True, sort = False)

        except FileNotFoundError as e:
            print(f'Arquivo não encontrado: {e}')        

    df_proc_nu_pf['data'] = df_nu_pf['date']
    df_proc_nu_pf['in_out'] = 'out'
    df_proc_nu_pf['categoria'] = 'cartão de crédito nu pf'
    df_proc_nu_pf['investimento'] = df_nu_pf['category']
    df_proc_nu_pf['quantidade'] = -100
    df_proc_nu_pf['preco'] = df_nu_pf['amount']
    df_proc_nu_pf['valor'] = df_nu_pf['amount']
    df_proc_nu_pf['status'] = 'pago'
    df_proc_nu_pf['descricao'] = df_nu_pf['title']

    valor_empty(path_base, df_proc_nu_pf)

    df_proc_nu_pf.to_csv(path_base + 'tabelas/' + 'processed/nu_pf_processed.csv', index = False, sep = ';')

def salary(path_base, folder_tabelas):

    df_dim_salario = pd.read_excel(path_base + folder_tabelas + 'dim_salario.xlsx')

    df_processed = df_dim_salario

    df_processed['id'] = df_processed.index

    df_processed['salario_bruto'] = df_processed['salario_bruto'].str.replace('.', '').str.replace(',', '.').str.replace('R', '').str.replace('$', '').str.replace(' ', '')
    df_processed['beneficios'] = df_processed['beneficios'].str.replace('.', '').str.replace(',', '.').str.replace('R', '').str.replace('$', '').str.replace(' ', '')
    df_processed['descontos'] = df_processed['descontos'].str.replace('.', '').str.replace(',', '.').str.replace('R', '').str.replace('$', '').str.replace(' ', '')
    df_processed['salario_liquido'] = df_processed['salario_liquido'].str.replace('.', '').str.replace(',', '.').str.replace('R', '').str.replace('$', '').str.replace(' ', '')

    df_processed.astype({'salario_bruto': 'float', 'beneficios': 'float', 'descontos': 'float', 'salario_liquido': 'float'})
    df_processed['data_fechamento'] = pd.to_datetime(df_processed['data_fechamento'], format = "%Y-%m-%d")
    # df_processed['empresa'] = df_processed['empresa']

    start_time = datetime.strptime("01/01/1900 00:00:00", "%d/%m/%Y %H:%M:%S")

    df_processed.to_csv(path_base + folder_tabelas + 'processed/dim_salario.csv', index = False, sep = ';')

def proc_b(path_base, folder_tabelas, folder, file):

    df_b = pd.read_csv(path_base + folder + file, delimiter = ',', dtype = 'str', encoding = 'UTF-8')
    # print(df_b.head())

    df_processed = pd.DataFrame()

    df_processed['created_timestamp'] = df_b['Date(UTC)']
    df_processed['executed_timestamp'] = df_b['Time']
    df_processed['par'] = df_b['Pair']
    df_split_0 = df_b['Pair'].str.split(pat='BRL', expand=True)
    df_processed['m'] = df_split_0[0]
    df_processed['tipo'] = df_b['Type']
    df_processed['qtd_ordem'] = df_b['Order Amount']
    df_processed['qtd_executed'] = df_b['Executed']

    for i in range(len(df_b)):
        df_split_1 = df_b.at[i, 'Executed'].split(str(df_processed.at[i, 'm']))[0]
        df_processed.at[i, 'qtd_comprada'] = df_split_1

    df_processed['preco_ordem'] = df_b['Order Price']
    df_processed['preco_medio_ordem'] = df_b['Average Price']
    df_split_2 = df_b['Trading total'].str.split(pat='BRL', expand=True)
    df_processed['preco_reais'] = df_split_2[0].astype(float).round(4)
    df_processed['status'] = df_b['Status']

    # print(df_processed.head())
    df_filtered = df_processed[df_processed['status'].isin(['FILLED'])]
    # print(df_filtered.head())

    return df_filtered

def b(path_base, folder_tabelas, folder, file, qtd_files, cat_name, proc_file_path):
    
    all_dfs = []
    df_join_anos = pd.DataFrame()
    folder = folder.split(',')
    file = file.split(',')

    for i in range(qtd_files):

        df_processed = proc_b(path_base, folder_tabelas, folder[i], file[i])
        all_dfs.append(df_processed)

    df_join_anos = pd.concat(all_dfs, ignore_index=True)

    df_renamed = pd.DataFrame()

    # df_renamed['data'] = pd.to_datetime(df_join_anos['executed_timestamp']).dt.strftime('%d/%m/%Y')
    df_renamed['data'] = pd.to_datetime(df_join_anos['executed_timestamp'], format='%Y-%m-%d %H:%M:%S').dt.strftime('%Y-%m-%d')

    # df_processed_2024['data'] = pd.to_datetime(df_processed_2024['data'], format = '%Y-%m-%d')

    df_renamed['in_out'] = 'in'
    df_renamed['categoria'] = cat_name
    df_renamed['investimento'] = df_join_anos['m']
    df_renamed['quantidade'] = df_join_anos['qtd_comprada']
    df_renamed['preco'] = df_join_anos['preco_medio_ordem']
    df_renamed['valor'] = df_join_anos['preco_reais']
    df_renamed['status'] = 'investido'
    df_renamed['descricao'] = ''

    # print(df_renamed.head())
    df_join_anos.to_csv(path_base + proc_file_path + 'b_join_anos.csv', index = False, sep = ';', encoding = 'UTF-8')
    # print(df_renamed.dtypes)

    # df_filtered = df_renamed[(pd.to_datetime(df_renamed['data'], format = '%d/%m/%Y').dt.year >= 2024) & (pd.to_datetime(df_renamed['data'], format = '%d/%m/%Y').dt.month >= 4)]
    df_filtered = df_renamed[(pd.to_datetime(df_renamed['data'], format = '%Y-%m-%d').dt.year >= 2024) & (pd.to_datetime(df_renamed['data'], format = '%Y-%m-%d').dt.month >= 4)]

    df_filtered.to_csv(path_base + proc_file_path + 'b_fato_gastos.csv', index = False, sep = ';', encoding = 'UTF-8')

def connect_to_postgres(host, database, user, password):

    try:
        conn = psycopg2.connect(host=host, dbname=database, user=user, password=password)
        return conn
    except psycopg2.OperationalError as e:
        print(f'Error connecting to PostgreSQL: {e}')
        raise

def delete_all_rows(conn, table_name, special_query = False, query = '', params = ''):

    try:
        cur = conn.cursor()
        if special_query:
            cur.execute(f'DELETE FROM {table_name} where {query}', params)
        else:
            cur.execute(f'DELETE FROM {table_name}')
        conn.commit()
    except psycopg2.Error as e:
        print(f'Error deleting rows from table {table_name}: {e}')
        raise

def insert_data_from_csv(conn, csv_file, table_name):

    try:
        cur = conn.cursor()

        # Read CSV header for column names (assuming header row exists)
        with open(csv_file, 'r', encoding = 'utf-8') as f:
            reader = csv.reader(f, delimiter=';')
            headers = next(reader)

        # Construct column placeholder string
        column_placeholders = ', '.join(['%s' for _ in headers])

        # Prepare INSERT query with dynamic column names
        insert_query = f'INSERT INTO {table_name} ({','.join(headers)}) VALUES ({column_placeholders})'

        # Read CSV data and insert row by row
        with open(csv_file, 'r', encoding = 'utf-8') as f:
            reader = csv.reader(f, delimiter=';')
            next(reader)  # Skip the header
            for row in reader:
                # Check for empty strings and replace them with None
                row = [None if value == '' else value for value in row]
                cur.execute(insert_query, row)
        conn.commit()

    except psycopg2.Error as e:
        print(f'Error inserting data from CSV file: {e}')
        raise

if __name__ == '__main__':

    load_dotenv()

    rasp_ssh_port = os.getenv('rasp_ssh_port')
    usernm_sftp = os.getenv('usernm_sftp')
    win_path_plan_gas = os.getenv('win_path_plan_gas')

    rasp_ip = os.getenv('rasp_ip')
    rasp_db_nm = os.getenv('rasp_db_nm')
    rasp_metab_user = os.getenv('rasp_metab_user')
    rasp_pwd = os.getenv('rasp_pwd')

    rasp_home_path = os.getenv('rasp_home_path')

    file_fat_pla_gas_con_anos = os.getenv('file_fat_pla_gas_con_anos')
    table_nm_fat_plan_gas = os.getenv('table_nm_fat_plan_gas')
    col_inv = os.getenv('col_inv')

    file_salary = os.getenv('file_salary')
    table_nm_salary = os.getenv('table_nm_salary')

    file_fat_notas = os.getenv('file_fat_notas')
    table_nm_fat_notas = os.getenv('table_nm_fat_notas')

    file_fat_atv_pro = os.getenv('file_fat_atv_pro')
    table_nm_fat_atv_pro = os.getenv('table_nm_fat_atv_pro')
    atv_pro_load_years = os.getenv('atv_pro_load_years')
    product_to_cat = os.getenv('product_to_cat')
    product_to_sector = os.getenv('product_to_sector')
  
    file_nu_pf = os.getenv('file_nu_pf')

    file_b = os.getenv('file_b')
    b_cat_name = os.getenv('b_cat_name')
    b_list_folder = os.getenv('b_list_folder')
    b_list_file = os.getenv('b_list_file')
    b_proc_file_path = os.getenv('b_proc_file_path')
    b_c = os.getenv('b_c')

    update_databases = os.getenv('update_databases', 'False').lower() in ['true', '1', 'yes', 'y', 't']
    update_files_sftp = os.getenv('update_files_sftp', 'False').lower() in ['true', '1', 'yes', 'y', 't']

    ####
    if os.getenv('env') == 'win':
        path_base = win_path_plan_gas + 'proj_plano_gastos\\'
        folder_tabelas = '\\tabelas\\'
        system_path = win_path_plan_gas + 'proj_plano_gastos\\tabelas\\processed\\'
    else:
        path_base = '/home/fabio/proj_plano_gastos/'
        folder_tabelas = 'tabelas/'
        system_path = os.getenv('rasp_table_processed')

    ####
    if update_files_sftp:
        sftp_move_files(rasp_ip, rasp_ssh_port, usernm_sftp, rasp_pwd, win_path_plan_gas, rasp_home_path)

    ####

    df_processed_2023 = fato_2023(path_base, folder_tabelas)
    df_processed_2024 = fato_2024(path_base, folder_tabelas)
    join_anos(path_base, folder_tabelas, df_processed_2023, df_processed_2024)
    print('Fato gastos processadas com sucesso')

    notas_fiscais_2024(path_base, folder_tabelas)
    print('Notas fiscais processadas com sucesso')

    atv_pro(path_base, folder_tabelas, product_to_cat, product_to_sector, atv_pro_load_years)
    print('Atv/pas processadas com sucesso')

    # nu_pf(path_base, folder_tabelas)
    # print('Nu pf processado com sucesso')

    salary(path_base, folder_tabelas)
    print('Salário processado com sucesso')

    b(path_base, folder_tabelas, b_list_folder, b_list_file, 3, b_cat_name, b_proc_file_path)
    print('B processed successfully')

    ####
    if update_databases:
        try:
            conn = connect_to_postgres(rasp_ip, rasp_db_nm, rasp_metab_user, rasp_pwd)
            cur = conn.cursor()
            ####
            # cur.execute(f'SELECT * FROM {table_name}')
            # # Fetch all data from the cursor (alternatively use fetchone() or fetchmany())
            # data = cur.fetchall()
            # print('Retrieved data:')
            # for row in data:
            #     print(row)  # Each row is a tuple containing column values
            ####
            delete_all_rows(conn, table_nm_fat_plan_gas)

            delete_all_rows(conn, table_nm_fat_notas)
            
            delete_all_rows(conn, table_nm_fat_atv_pro)
            
            # delete_all_rows(conn, table_nm_fat_plan_gas, True, 'quantidade = -100')

            delete_all_rows(conn, table_nm_salary)

            delete_all_rows(conn, table_nm_fat_plan_gas, True, f'data = %s AND {col_inv} = %s', ('02/04/2024', b_c))

            ####
            insert_data_from_csv(conn, system_path + file_fat_pla_gas_con_anos, table_nm_fat_plan_gas)
            print('Data successfully imported from CSV to Fato Plano Gastos table.')

            insert_data_from_csv(conn, system_path + file_fat_notas, table_nm_fat_notas)
            print('Data successfully imported from CSV to Fato Notas table.')

            insert_data_from_csv(conn, system_path + file_fat_atv_pro, table_nm_fat_atv_pro)
            print('Data successfully imported from CSV to Fato Atv Pas table.')

            # insert_data_from_csv(conn, system_path + file_nu_pf, table_nm_fat_plan_gas)
            # print('Data successfully imported from CSV to Fato Nu pf table.')

            insert_data_from_csv(conn, system_path + file_salary, table_nm_salary)
            print('Data successfully imported from CSV to Salary table.')

            insert_data_from_csv(conn, path_base + b_proc_file_path + file_b, table_nm_fat_plan_gas)
            print('Data successfully imported from CSV to B table.')
            ####
        except Exception as e:
            print(f'An error occurred: {e}')
        finally:
            if conn:
                conn.close()
                print('Connection to PostgreSQL closed.')